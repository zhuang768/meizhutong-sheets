#!/usr/bin/env python3
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

SCENARIO_ZH = {
    'pass': '合格：一般青年',
    'passSpecial': '合格：特定對象',
    'passCultural': '合格：文化語言保存者',
    'passMonthly': '合格：月費制',
    'passProxy': '合格：他人代付',
    'passXiangshan': '合格：香山區',
    'passSyntheticId': '合格：合成身分證',
    'passHighAmount': '合格：高額達補助上限',
    'notResident': '未設籍新竹市',
    'prepaid': '預付點數／Token',
    'aggregator': '集合式平台購買',
    'reseller': '代購非官方',
    'vendorRegion': '廠商地區不符',
    'ageTooOld': '年齡超過上限',
    'ageTooYoung': '年齡低於下限',
    'purchaseTooEarly': '購買日過早',
    'purchaseTooLate': '購買日過晚',
    'applyLateMonthly': '月費逾期申請',
    'badId': '身分證檢查碼錯誤',
    'missingId': '未填身分證字號',
    'missingName': '未填姓名',
    'missingDocs': '缺所有應備文件',
    'missingReceipt': '未勾選收據確認',
    'specialMissing': '特定對象缺證明',
    'culturalMissing': '文化語言缺證照',
    'proxyMissing': '代付缺證明文件',
    'proxyMissingName': '代付未填姓名關係',
    'missingBank': '未填銀行帳戶',
    'badContact': '聯絡資料格式錯誤',
    'zeroAmount': '金額為 0',
    'periodInverted': '訂閱迄日早於起日',
    'missingChannel': '未填購買管道',
    'missingTool': '未填工具或廠商',
    'addressInconsistent': '設籍與地址不一致',
    'badDistrict': '行政區不是新竹市',
    'badPostal': '郵遞區號不是 300',
    'specialNoKind': '特定對象未勾身分類別',
    'missingBirth': '未填出生日期',
    'missingPassbook': '缺存摺封面',
    'duplicate': '重複身分證字號',
}

def header_row(ws, values, fill, font):
    cells = []
    for value in values:
        cell = WriteOnlyCell(ws, value=value)
        cell.fill = fill
        cell.font = font
        cells.append(cell)
    ws.append(cells)

def export_xlsx(count, csv_path, jsonl_path, summary_path, out_path):
    summary = json.loads(Path(summary_path).read_text())
    header_fill = PatternFill('solid', fgColor='1F4E79')
    header_font = Font(bold=True, color='FFFFFF')
    true_fill = PatternFill('solid', fgColor='D4E2D1')
    false_fill = PatternFill('solid', fgColor='E6D3D1')

    def paint(ws, value):
        if value is True or value == 'TRUE':
            cell = WriteOnlyCell(ws, value='TRUE')
            cell.fill = true_fill
            return cell
        if value is False or value == 'FALSE':
            cell = WriteOnlyCell(ws, value='FALSE')
            cell.fill = false_fill
            return cell
        text = str(value)
        if text == '需人工複核' or text.startswith('建議駁回'):
            cell = WriteOnlyCell(ws, value=value)
            cell.fill = false_fill
            return cell
        if text.startswith('建議核准'):
            cell = WriteOnlyCell(ws, value=value)
            cell.fill = true_fill
            return cell
        if text == '建議補件':
            cell = WriteOnlyCell(ws, value=value)
            cell.fill = PatternFill('solid', fgColor='F3E6C8')
            return cell
        return value
    wb = Workbook(write_only=True)

    info = wb.create_sheet('說明')
    info.column_dimensions['A'].width = 22
    info.column_dimensions['B'].width = 72
    info.append(['梅竹通合成測試資料'])
    info.append(['筆數', count])
    info.append(['用途', '給承辦／AI 查核測試，非正式申請'])
    info.append(['人工審查決定', '全部為待審，AI 不可自動核准或駁回'])
    info.append(['身分證等個資', '皆為合成資料，非正式證件'])
    info.append(['規則建議', '見「統計」與「合成案件」前幾欄；ChatGPT 尚未寫入'])
    info.append(['產生時間', summary.get('generatedAt', '')])

    stats = wb.create_sheet('統計')
    stats.column_dimensions['A'].width = 28
    stats.column_dimensions['B'].width = 22
    stats.column_dimensions['C'].width = 14
    header_row(stats, ['規則建議', '筆數'], header_fill, header_font)
    for label, value in summary['bySuggestion'].items():
        stats.append([label, value])
    stats.append([])
    header_row(stats, ['信心程度', '筆數'], header_fill, header_font)
    for label, value in summary['byConfidence'].items():
        stats.append([label, value])
    stats.append([])
    header_row(stats, ['情境代碼', '中文說明', '筆數'], header_fill, header_font)
    for key in summary['scenarios']:
        stats.append([key, SCENARIO_ZH.get(key, key), summary['byScenario'][key]])

    cases = wb.create_sheet('合成案件')
    cases.freeze_panes = 'A2'
    cases.column_dimensions['A'].width = 18
    cases.column_dimensions['B'].width = 22
    cases.column_dimensions['C'].width = 22
    cases.column_dimensions['D'].width = 10
    cases.column_dimensions['E'].width = 20

    with Path(csv_path).open('r', encoding='utf-8-sig', newline='') as csv_file, Path(jsonl_path).open('r', encoding='utf-8') as jsonl_file:
        reader = csv.reader(csv_file)
        headers = next(reader)
        header_row(cases, ['情境代碼', '合成情境', '規則建議', '規則信心'] + headers, header_fill, header_font)
        for row, line in zip(reader, jsonl_file):
            rec = json.loads(line)
            expected = rec.get('expected') or {}
            scenario = rec.get('scenario', '')
            cases.append([paint(cases, value) for value in [
                scenario,
                SCENARIO_ZH.get(scenario, scenario),
                expected.get('suggestion', ''),
                expected.get('confidence', ''),
                *row,
            ]])

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

if __name__ == '__main__':
    count = int(sys.argv[1]) if len(sys.argv) > 1 else 50000
    root = Path(__file__).resolve().parents[2]
    testdata = root / 'testdata'
    default_out = Path.home() / 'Desktop' / f'梅竹通-合成測試-{count}.xlsx'
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else default_out
    export_xlsx(
        count,
        testdata / f'loadtest-{count}.csv',
        testdata / f'loadtest-{count}.jsonl',
        testdata / f'loadtest-{count}.summary.json',
        out_path,
    )
    print(out_path)
    print(out_path.stat().st_size)
